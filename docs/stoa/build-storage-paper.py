# -*- coding: utf-8 -*-
"""Markdown -> .docx for the StoaChain storage / Arweave-offload paper.

House style matches docs/stoa/_private/build-xchain-paper.py (Calibri 10.5,
"Light Grid Accent 1" tables, AncientHoldings GmbH core properties).

Usage:
    py docs/stoa/build-storage-paper.py                 # default in/out
    py docs/stoa/build-storage-paper.py IN.md OUT.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

C_GREY = RGBColor(0x70, 0x70, 0x70)
C_BLUE = RGBColor(0x1F, 0x4E, 0x79)
C_CODE = RGBColor(0x1A, 0x1A, 0x1A)
SHADE_CODE = "F2F2F2"

HERE = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "storage-architecture-and-arweave-offload.md"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "StoaChain-Storage-Architecture-and-Arweave-Offload.docx"

doc = Document()
_n = doc.styles["Normal"]
_n.font.name = "Calibri"
_n.font.size = Pt(10.5)

# A4, since the audience reads in Europe.
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.9)

# ---------------------------------------------------------------- inline spans

INLINE = re.compile(
    r"(\*\*.+?\*\*)"            # bold
    r"|(`[^`]+`)"               # code
    r"|(\[[^\]]*\]\([^)]*\))"   # link
    r"|(\*[^*\n]+\*)"           # italic
)


def add_inline(p, text, base_size=None, base_color=None):
    """Append `text` to paragraph `p`, honouring bold/italic/code/link markup."""
    def run(s, **kw):
        r = p.add_run(s)
        if base_size is not None:
            r.font.size = Pt(base_size)
        if base_color is not None:
            r.font.color.rgb = base_color
        for k, v in kw.items():
            if k == "color":
                r.font.color.rgb = v
            elif k == "name":
                r.font.name = v
            elif k == "size":
                r.font.size = Pt(v)
            else:
                setattr(r, k, v)
        return r

    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            run(tok[2:-2], bold=True)
        elif tok.startswith("`"):
            run(tok[1:-1], name="Consolas",
                size=(base_size or 10.5) - 1.0, color=C_CODE)
        elif tok.startswith("["):
            label, url = re.match(r"\[([^\]]*)\]\(([^)]*)\)", tok).groups()
            # The label may itself be `code`.
            inner = label[1:-1] if label.startswith("`") and label.endswith("`") else label
            run(inner, color=C_BLUE,
                name="Consolas" if label.startswith("`") else None)
            if url.startswith("http"):
                run(" (" + url + ")", color=C_GREY, size=(base_size or 10.5) - 2.0)
        else:
            run(tok[1:-1], italic=True)
        pos = m.end()
    if pos < len(text):
        run(text[pos:])


# ------------------------------------------------------------------- block ops

def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pPr.append(sh)


def add_code_block(lines):
    for ln in lines:
        p = doc.add_paragraph()
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = C_CODE
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Inches(0.18)
        shade(p, SHADE_CODE)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    """Split a markdown table row on unescaped pipes."""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    s = s.replace("\\|", "\x00")
    return [c.strip().replace("\x00", "|") for c in s.split("|")]


def add_table(rows):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            val = row[j] if j < len(row) else ""
            if i == 0:
                add_inline(p, "**" + val + "**" if val and "**" not in val else val,
                           base_size=8.5)
            else:
                add_inline(p, val, base_size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# ----------------------------------------------------------------- the parser

md = SRC.read_text(encoding="utf-8").splitlines()

# Core properties, matching the house convention.
cp = doc.core_properties
cp.author = "AncientHoldings GmbH"
cp.last_modified_by = "AncientHoldings GmbH"
cp.title = "StoaChain Storage Architecture and the Arweave Offload Plan"
cp.subject = ("How a chainweb node stores data, why only half of it can be "
              "offloaded, and what it would take to put block history on Arweave")
cp.category = "Design Paper"
cp.version = "1.0"
cp.keywords = "StoaChain, chainweb, pact, RocksDB, Arweave, storage, AncientHoldings GmbH"

i = 0
first_heading_done = False
in_front_matter = False

while i < len(md):
    line = md[i]
    stripped = line.strip()

    # fenced code
    if stripped.startswith("```"):
        i += 1
        buf = []
        while i < len(md) and not md[i].strip().startswith("```"):
            buf.append(md[i])
            i += 1
        i += 1
        add_code_block(buf)
        continue

    # tables
    if stripped.startswith("|") and i + 1 < len(md) and re.match(r"^\|[\s:|-]+\|?$", md[i + 1].strip()):
        rows = [split_row(stripped)]
        i += 2
        while i < len(md) and md[i].strip().startswith("|"):
            rows.append(split_row(md[i].strip()))
            i += 1
        add_table(rows)
        continue

    # headings
    m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if m:
        level, text = len(m.group(1)), m.group(2)
        if level == 1 and not first_heading_done:
            doc.add_heading(text, level=0)
            first_heading_done = True
            in_front_matter = True
        else:
            in_front_matter = False
            h = doc.add_heading("", level=min(level - 1, 4) if level > 1 else 1)
            add_inline(h, text)
        i += 1
        continue

    # horizontal rule -> thin spacer
    if stripped in ("---", "***", "___"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        i += 1
        continue

    # blank
    if not stripped:
        i += 1
        continue

    # block quote
    if stripped.startswith(">"):
        buf = []
        while i < len(md) and md[i].strip().startswith(">"):
            buf.append(md[i].strip().lstrip(">").strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))
        for r in p.runs:
            r.italic = True
        pf = p.paragraph_format
        pf.left_indent = Inches(0.3)
        pf.space_after = Pt(6)
        continue

    # bullets / numbered
    mb = re.match(r"^[-*]\s+(.*)$", stripped)
    mn = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
    if mb or mn:
        text = mb.group(1) if mb else mn.group(2)
        style = "List Bullet" if mb else "List Number"
        i += 1
        # Gather the item's continuation lines BEFORE rendering, so an inline
        # span that wraps across a line break is still matched as one span.
        parts = [text]
        while i < len(md) and md[i].startswith("  ") and md[i].strip() \
                and not re.match(r"^\s*([-*]|\d+[.)])\s", md[i]) \
                and not md[i].strip().startswith("|"):
            parts.append(md[i].strip())
            i += 1
        p = doc.add_paragraph(style=style)
        add_inline(p, " ".join(parts))
        p.paragraph_format.space_after = Pt(3)
        continue

    # ordinary paragraph: gather until blank / block start
    buf = [stripped]
    i += 1
    while i < len(md):
        nxt = md[i].strip()
        if not nxt or nxt.startswith(("#", "|", "```", ">", "---")) \
                or re.match(r"^([-*]|\d+[.)])\s", nxt):
            break
        buf.append(nxt)
        i += 1
    p = doc.add_paragraph()
    add_inline(p, " ".join(buf))
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    if in_front_matter:
        for r in p.runs:
            if not r.bold:
                r.font.color.rgb = C_GREY

# byline, inserted just after the title
byline = doc.paragraphs[1] if len(doc.paragraphs) > 1 else doc.add_paragraph()
bp = byline.insert_paragraph_before()
r = bp.add_run("Prepared by AncientHoldings GmbH")
r.bold = True
r.font.size = Pt(11)
bp.paragraph_format.space_after = Pt(1)
bp2 = byline.insert_paragraph_before()
r2 = bp2.add_run("A product of AncientHoldings GmbH · 2026-10-23 · Version 1.0")
r2.font.color.rgb = C_GREY
r2.font.size = Pt(10)
bp2.paragraph_format.space_after = Pt(10)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print("wrote", OUT)
print("paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables))
